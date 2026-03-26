import openpyxl

def extract_excel_info(filename):
    wb = openpyxl.load_workbook(filename, data_only=False)
    wb_data = openpyxl.load_workbook(filename, data_only=True)
    
    sheet_name = "Cálculos de escuadrado"
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name} not found. Available: {wb.sheetnames}")
        return

    sheet = wb[sheet_name]
    sheet_data = wb_data[sheet_name]
    
    print(f"--- Logic from {sheet_name} ---")
    for row in range(1, 40):
        row_vals = []
        for col in range(1, 15):
            cell = sheet.cell(row=row, column=col)
            cell_data = sheet_data.cell(row=row, column=col)
            val = cell_data.value
            formula = cell.value if isinstance(cell.value, str) and cell.value.startswith('=') else None
            
            if val is not None or formula is not None:
                coord = f"{cell.column_letter}{cell.row}"
                content = f"{coord}: {val}"
                if formula:
                    content += f" (F: {formula})"
                row_vals.append(content)
        if row_vals:
            print(" | ".join(row_vals))

if __name__ == "__main__":
    extract_excel_info("Calculo Iris.xlsx")
